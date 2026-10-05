import os
import sqlite3
import openpyxl
from flask import Flask, render_template, request, jsonify

app = Flask(__name__)
DB_NAME = 'vedveda.db'
EXCEL_FILE = 'VEDVEDA.xlsx'

def get_db_connection():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    conn = get_db_connection()
    cursor = conn.cursor()
    
    # Create tables
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS members (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            role TEXT DEFAULT 'Team Member'
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            member_id INTEGER,
            date TEXT,
            product_name TEXT,
            quantity INTEGER DEFAULT 1,
            pv REAL DEFAULT 0,
            amount REAL DEFAULT 0,
            payment_status TEXT DEFAULT 'Pending',
            penal_pv REAL DEFAULT 0,
            penal_amount REAL DEFAULT 0,
            FOREIGN KEY (member_id) REFERENCES members (id)
        )
    ''')
    conn.commit()
    
    # Seed from Excel if transactions table is empty
    cursor.execute('SELECT COUNT(*) FROM transactions')
    count = cursor.fetchone()[0]
    if count == 0 and os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
            
            # Map canonical team members
            members_map = {
                "Pooja": "Team Member",
                "Deepika": "Team Member",
                "Pankaj": "Team Member",
                "Shailendra": "Team Member",
                "Anita": "Team Member"
            }
            
            for m_name, role in members_map.items():
                cursor.execute('INSERT OR IGNORE INTO members (name, role) VALUES (?, ?)', (m_name, role))
            conn.commit()
            
            # Helper dictionary for member IDs
            cursor.execute('SELECT id, name FROM members')
            m_ids = {row['name']: row['id'] for row in cursor.fetchall()}
            
            # Sheet mapping
            sheets_to_parse = {
                "POOJA": "Pooja",
                "SHAILE": "Shailendra",
                "PANKAJ": "Pankaj",
                "Anita mem": "Anita",
                "depika mem": "Deepika"
            }
            
            for sheet_name, member_name in sheets_to_parse.items():
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    member_id = m_ids.get(member_name)
                    if not member_id:
                        continue
                    
                    # Find header row
                    header_row = None
                    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
                        row_strs = [str(c).strip().upper() if c else "" for c in row]
                        if "DATE" in row_strs:
                            header_row = i
                            break
                    
                    if header_row:
                        last_date = None
                        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                            raw_date = row[0]
                            raw_product = row[1] if len(row) > 1 else None
                            raw_pv = row[2] if len(row) > 2 else None
                            raw_amt = row[3] if len(row) > 3 else None
                            raw_payment = row[4] if len(row) > 4 else None
                            raw_penal = row[5] if len(row) > 5 else None
                            
                            if all(v is None for v in [raw_date, raw_product, raw_pv, raw_amt]):
                                continue
                            
                            product = str(raw_product).strip().title() if raw_product else ""
                            if not product or product.lower() in ("total", "no", "product"):
                                continue
                            
                            if raw_date:
                                last_date = str(raw_date).split()[0]
                            
                            try:
                                pv = float(raw_pv or 0)
                            except ValueError:
                                pv = 0.0
                            
                            try:
                                amt = float(raw_amt or 0)
                            except ValueError:
                                amt = 0.0
                            
                            paid_val = float(raw_payment or 0) if raw_payment else 0.0
                            if paid_val >= amt and amt > 0:
                                status = "Paid"
                                penal_amt = 0.0
                            elif 0 < paid_val < amt:
                                status = "Pending"
                                penal_amt = round(amt - paid_val, 2)
                            else:
                                status = "Pending" if amt > 0 else "Paid"
                                penal_amt = amt if amt > 0 else 0.0
                            
                            cursor.execute('''
                                INSERT INTO transactions (member_id, date, product_name, pv, amount, payment_status, penal_amount)
                                VALUES (?, ?, ?, ?, ?, ?, ?)
                            ''', (member_id, last_date, product, pv, amt, status, penal_amt))
            
            conn.commit()
            print("Successfully seeded database from Excel using openpyxl.")
        except Exception as e:
            print(f"Error seeding from Excel: {e}")
            
    conn.close()

@app.route('/')
def index():
    init_db()
    return render_template('dashboard.html')

@app.route('/api/stats/overview')
def get_stats_overview():
    init_db()
    conn = get_db_connection()
    cursor = conn.cursor()
    
    cursor.execute('SELECT m.name, SUM(t.pv) as pv FROM members m LEFT JOIN transactions t ON m.id = t.member_id GROUP BY m.id')
    pv_per_member = [{"member": row["name"], "pv": round(row["pv"] or 0, 2)} for row in cursor.fetchall()]
    
    cursor.execute('SELECT SUM(CASE WHEN payment_status="Paid" THEN amount ELSE 0 END) as paid, SUM(CASE WHEN payment_status="Pending" THEN amount ELSE 0 END) as pending FROM transactions')
    res = cursor.fetchone()
    
    conn.close()
    return jsonify({
        "pv_per_member": pv_per_member,
        "paid_vs_pending": {
            "paid": round(res["paid"] or 0, 2),
            "pending": round(res["pending"] or 0, 2)
        }
    })

@app.route('/api/members')
def api_members():
    init_db()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT * FROM members ORDER BY name')
    members = []
    for row in cursor.fetchall():
        m_id = row['id']
        cursor.execute('SELECT SUM(pv) as tpv, SUM(amount) as tamt, SUM(CASE WHEN payment_status="Paid" THEN amount ELSE 0 END) as pamt, SUM(CASE WHEN payment_status="Pending" THEN amount ELSE 0 END) as pamtd, COUNT(*) as cnt FROM transactions WHERE member_id = ?', (m_id,))
        stats = cursor.fetchone()
        members.append({
            "id": m_id,
            "name": row['name'],
            "role": row['role'],
            "total_pv": round(stats['tpv'] or 0, 2),
            "total_amount": round(stats['tamt'] or 0, 2),
            "paid_amount": round(stats['pamt'] or 0, 2),
            "pending_amount": round(stats['pamtd'] or 0, 2),
            "transaction_count": stats['cnt'] or 0
        })
    conn.close()
    return jsonify(members)

@app.route('/api/transactions')
def api_transactions():
    init_db()
    conn = get_db_connection()
    cursor = conn.cursor()
    member_id = request.args.get('member_id', type=int)
    
    query = '''
        SELECT t.*, m.name as member_name 
        FROM transactions t 
        JOIN members m ON t.member_id = m.id
    '''
    params = []
    if member_id:
        query += ' WHERE t.member_id = ?'
        params.append(member_id)
    query += ' ORDER BY t.date DESC'
    
    cursor.execute(query, params)
    txns = [dict(row) for row in cursor.fetchall()]
    conn.close()
    return jsonify(txns)

if __name__ == '__main__':
    init_db()
    app.run(host='0.0.0.0', port=5000, debug=True)